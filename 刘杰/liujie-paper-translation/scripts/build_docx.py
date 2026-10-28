#!/usr/bin/env python3
"""Build an editable Chinese paper translation from the skill's block JSON.

Requires python-docx, Pillow and lxml. Native math additionally needs
latex2mathml and a locally supplied MML2OMML.XSL; no Office file is bundled.
This builder is not a translator, completeness auditor, or visual QA tool.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image


KINDS = {"heading", "paragraph", "figure", "table", "equation", "reference", "list_item", "note"}


def text(value, label, *, allow_empty=False):
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError(f"{label} must be {'a' if allow_empty else 'a nonempty'} string")
    return value


def scalar(value, label):
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ValueError(f"{label} must be a string or finite number, not a structured object")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{label} must be finite")
    return str(value)


def element(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    return node


def validate(data):
    if not isinstance(data, dict) or data.get("schema_version") != 1:
        raise ValueError("Input must be a schema_version=1 JSON object")
    text(data.get("title_zh"), "title_zh")
    text(data.get("title_en"), "title_en", allow_empty=True)
    if "metadata_line" in data:
        text(data["metadata_line"], "metadata_line", allow_empty=True)
    sources = data.get("sources")
    blocks = data.get("blocks")
    if not isinstance(sources, list) or not isinstance(blocks, list) or not blocks:
        raise ValueError("sources must be a list; blocks must be a nonempty list")
    source_ids = []
    for source in sources:
        if not isinstance(source, dict):
            raise ValueError("Every source must be an object")
        source_ids.append(text(source.get("id"), "source.id"))
    if len(set(source_ids)) != len(source_ids):
        raise ValueError("Duplicate source IDs")
    front = data.get("front_matter_source_ids", {})
    if not isinstance(front, dict) or set(front) - {"title_zh", "title_en", "metadata_line"}:
        raise ValueError("front_matter_source_ids must map title_zh/title_en/metadata_line to source ID lists")
    for field, ids in front.items():
        if not isinstance(ids, list) or any(not isinstance(x, str) for x in ids):
            raise ValueError(f"front_matter_source_ids.{field} must be an explicit list of IDs")
        if len(set(ids)) != len(ids) or set(ids) - set(source_ids):
            raise ValueError(f"front_matter_source_ids.{field}: duplicate or unknown source IDs")
        if ids and not data.get(field):
            raise ValueError(f"front_matter_source_ids.{field}: cannot map source content to an absent or empty rendered field")
    seen = set()
    for block in blocks:
        if not isinstance(block, dict):
            raise ValueError("Every block must be an object")
        bid = text(block.get("id"), "block.id")
        if bid in seen:
            raise ValueError(f"Duplicate block ID: {bid}")
        seen.add(bid)
        if block.get("kind") not in KINDS:
            raise ValueError(f"{bid}: unsupported kind {block.get('kind')!r}")
        ids = block.get("source_ids")
        if not isinstance(ids, list) or not ids or any(not isinstance(x, str) for x in ids):
            raise ValueError(f"{bid}: source_ids must be an explicit nonempty list of IDs")
        if len(set(ids)) != len(ids) or set(ids) - set(source_ids):
            raise ValueError(f"{bid}: duplicate or unknown source_ids")
        kind = block["kind"]
        if kind not in {"equation", "reference"}:
            text(block.get("zh"), f"{bid}.zh")
        if kind == "reference":
            text(block.get("text"), f"{bid}.text")
        if kind == "heading":
            if type(block.get("level", 1)) is not int or block.get("level", 1) not in (1, 2, 3):
                raise ValueError(f"{bid}: heading level must be 1, 2 or 3")
        if "number" in block and (type(block["number"]) not in (str, int) or not str(block["number"]).strip()):
            raise ValueError(f"{bid}: number must be a nonempty string or integer")
        if kind in {"figure", "table"}:
            if not scalar(block.get("number"), f"{bid}.number").strip():
                raise ValueError(f"{bid}: number must be nonempty")
        if kind == "equation":
            if bool(block.get("latex")) == bool(block.get("image")):
                raise ValueError(f"{bid}: supply exactly one of latex or image")
            if block.get("latex"):
                text(block["latex"], f"{bid}.latex")
        if kind == "table":
            headers, rows = block.get("headers"), block.get("rows")
            if "allow_row_split" in block and type(block["allow_row_split"]) is not bool:
                raise ValueError(f"{bid}.allow_row_split must be a boolean")
            if not isinstance(headers, list) or not headers or not isinstance(rows, list):
                raise ValueError(f"{bid}: table requires nonempty headers and rows array")
            for ri, row in enumerate([headers] + rows):
                if not isinstance(row, list) or len(row) != len(headers):
                    raise ValueError(f"{bid}: row {ri} width differs from header width")
                for ci, value in enumerate(row):
                    scalar(value, f"{bid}.cell[{ri},{ci}]")
            if "notes" in block:
                text(block["notes"], f"{bid}.notes", allow_empty=True)
        if "labels_zh" in block:
            labels = block["labels_zh"]
            if not isinstance(labels, dict):
                raise ValueError(f"{bid}.labels_zh must map English labels to Chinese strings")
            for key, value in labels.items():
                text(key, f"{bid}.labels_zh key")
                text(value, f"{bid}.labels_zh[{key}]")


class Builder:
    def __init__(self, source, args):
        self.source, self.args = source, args
        self.doc = Document()
        self.transform = None
        self.math_count = 0
        self.image_equations = 0
        self.bookmarks = []
        self.asset_paths = set()
        section = self.doc.sections[0]
        width, height = (8.5, 11) if args.page_size == "letter" else (210 / 25.4, 297 / 25.4)
        section.page_width, section.page_height = Inches(width), Inches(height)
        section.top_margin, section.bottom_margin = Inches(.72), Inches(.7)
        section.left_margin = section.right_margin = Inches(.8)
        section.header_distance, section.footer_distance = Inches(.25), Inches(.28)
        self.content_width = width - 1.6
        sizes = {"Normal": 11, "Title": 18, "Heading 1": 14, "Heading 2": 12, "Heading 3": 11, "Caption": 10}
        for name, size in sizes.items():
            style = self.doc.styles[name]
            style.font.name, style.font.size = args.font_latin, Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), args.font_heading_cjk if name.startswith("Heading") or name == "Title" else args.font_cjk)
            fmt = style.paragraph_format
            fmt.line_spacing = 1.18
            fmt.space_after = Pt(5)
            fmt.space_before = Pt(10 if name.startswith("Heading") else 0)
            fmt.keep_with_next = name.startswith("Heading")
            fmt.widow_control = True
        for style in self.doc.styles:
            for border in list(style._element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
            for color in style._element.iter(qn("w:color")):
                for attr in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn("w:" + attr), None)
        self.doc.settings.element.append(element("w:updateFields", val="true"))
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = element("w:fldSimple", instr="PAGE")
        cached = footer.add_run("1")
        self.fonts(cached, 9)
        field.append(cached._r)
        footer._p.append(field)
        self.doc.core_properties.author = ""
        self.doc.core_properties.last_modified_by = ""

    def fonts(self, run, size=11, *, bold=False, code=False, heading=False):
        run.font.name = "Consolas" if code else self.args.font_latin
        run.font.size, run.font.color.rgb, run.bold = Pt(size), RGBColor(0, 0, 0), bold
        rf = run._element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), self.args.font_heading_cjk if heading else self.args.font_cjk)

    def load_transform(self):
        if self.transform is not None:
            return
        supplied = self.args.mml2omml or os.environ.get("MML2OMML_XSL")
        candidates = [Path(supplied).expanduser()] if supplied else []
        if not supplied:
            for env in ("ProgramFiles", "ProgramFiles(x86)"):
                root = os.environ.get(env)
                if root:
                    for folder in ("Microsoft Office/root/Office16", "Microsoft Office/Office16", "Microsoft Office/Office15"):
                        candidates.append(Path(root) / folder / "MML2OMML.XSL")
        stylesheet = next((p for p in candidates if p.is_file()), None)
        if stylesheet is None:
            raise ValueError("Native math requires a local MML2OMML.XSL. Pass --mml2omml PATH or set MML2OMML_XSL. Do not substitute plain text or an image without approval.")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        self.transform = etree.XSLT(etree.parse(str(stylesheet), parser), access_control=etree.XSLTAccessControl.DENY_ALL)

    def add_math(self, paragraph, latex, size=11):
        if any(char in latex for char in ("\t", "\f", "\b")):
            raise ValueError("Equation contains a tab, form-feed or backspace; check JSON escaping (LaTeX backslashes must be doubled).")
        self.load_transform()
        try:
            from latex2mathml.converter import convert
        except ImportError as exc:
            raise ValueError("Native math requires latex2mathml in the selected runtime") from exc
        try:
            root = self.transform(etree.fromstring(convert(latex).encode("utf-8"))).getroot()
            if root is None or root.tag not in {qn("m:oMath"), qn("m:oMathPara")}:
                raise ValueError("converter returned no Office Math element")
            for math_run in root.iter(qn("m:r")):
                props = element("w:rPr")
                props.append(element("w:rFonts", ascii="Cambria Math", hAnsi="Cambria Math"))
                props.append(element("w:sz", val=int(size * 2)))
                math_run.insert(0, props)
            paragraph._p.append(root)
            self.math_count += 1
        except Exception as exc:
            raise ValueError(f"Cannot convert equation {latex!r}: {exc}") from exc

    def rich(self, paragraph, value, size=11, *, bold=False, heading=False):
        value = text(value, "rich text", allow_empty=True).replace("\u00ad", "")
        # Escaped literal dollars are not math delimiters. Formatting tokens are
        # deliberately small: inline $math$, **bold**, and `code`, not Markdown.
        literal_dollar = "\ue000"
        if literal_dollar in value:
            raise ValueError("Reserved private-use character in text")
        value = value.replace(r"\$", literal_dollar)
        for part in re.split(r"(\$[^$]+\$|\*\*[^*]+\*\*|`[^`]+`)", value):
            if not part:
                continue
            if part.startswith("$") and part.endswith("$"):
                self.add_math(paragraph, part[1:-1], size)
            else:
                if "$" in part:
                    raise ValueError("Unbalanced $ math delimiter; escape a literal dollar as \\$")
                is_bold = part.startswith("**") and part.endswith("**")
                is_code = part.startswith("`") and part.endswith("`")
                content = part[2:-2] if is_bold else part[1:-1] if is_code else part
                run = paragraph.add_run(content.replace(literal_dollar, "$"))
                self.fonts(run, size, bold=bold or is_bold, code=is_code, heading=heading)

    def bookmark(self, paragraph, block, index):
        name = f"B{index:05d}_" + re.sub(r"[^A-Za-z0-9_]", "_", block["id"])
        name = name[:40]
        start = element("w:bookmarkStart", id=index, name=name)
        end = element("w:bookmarkEnd", id=index)
        # Keep paragraph properties first, as required by WordprocessingML.
        pos = 1 if paragraph._p.pPr is not None else 0
        paragraph._p.insert(pos, start)
        paragraph._p.append(end)
        self.bookmarks.append({"block_id": block["id"], "source_ids": block["source_ids"], "bookmark": name})

    def image(self, block, *, equation=False):
        raw = Path(text(block.get("image"), f"{block['id']}.image"))
        if raw.is_absolute() or raw.drive:
            raise ValueError(f"{block['id']}: image paths must be relative to the JSON file")
        path = (self.source.parent / raw).resolve()
        if not path.is_file():
            raise ValueError(f"Image not found: {raw}")
        self.asset_paths.add(path)
        with Image.open(path) as image:
            pixel_width, pixel_height = image.size
        width = min(self.content_width, 2.5 * pixel_width / pixel_height) if equation else min(self.content_width, 5 * pixel_width / pixel_height)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = not equation
        shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
        shape._inline.docPr.set("descr", text(block.get("alt", block.get("zh", "")), "alt", allow_empty=True))
        return paragraph

    @staticmethod
    def numbered(label, number, value):
        number = scalar(number, "number")
        return value if re.match(rf"^{label}\s*{re.escape(number)}(?![\d.])", value) else f"{label} {number}  {value}"

    def add_table(self, block):
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = True
        self.rich(caption, self.numbered("表", block["number"], block["zh"]), 10)
        headers, rows = block["headers"], block["rows"]
        count = len(headers)
        widths = block.get("column_widths_inches")
        if widths is not None:
            if not isinstance(widths, list) or len(widths) != count or any(type(w) not in (int, float) or not math.isfinite(w) or w <= 0 for w in widths):
                raise ValueError(f"{block['id']}: widths must be one positive number per column")
            if sum(widths) > self.content_width + .01:
                raise ValueError(f"{block['id']}: column widths exceed content width {self.content_width:.2f}in")
        else:
            # Content-based starting point; visual QA still decides final widths.
            weights = [math.sqrt(max(6, min(40, max(len(str(row[i])) for row in [headers] + rows)))) for i in range(count)]
            widths = [self.content_width * weight / sum(weights) for weight in weights]
        table = self.doc.add_table(rows=1, cols=count)
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        borders = element("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(element("w:" + side, val="single", sz="4", color="D9D9D9"))
        table._tbl.tblPr.append(borders)
        for ri, values in enumerate([headers] + rows):
            row = table.rows[0] if ri == 0 else table.add_row()
            props = row._tr.get_or_add_trPr()
            if ri == 0:
                props.append(element("w:tblHeader"))
            if ri == 0 or not block.get("allow_row_split", False):
                props.append(element("w:cantSplit"))
            for ci, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Inches(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                margins = element("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margins.append(element("w:" + side, w=65 if side in ("top", "bottom") else 70, type="dxa"))
                cp.append(margins)
                if ri == 0:
                    cp.append(element("w:shd", fill="EAF0F5"))
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell_text = scalar(value, "table cell")
                numeric_text = re.sub(r"\\(?:pm|times|cdot)", "±", cell_text.strip("$"))
                numeric = re.fullmatch(r"[0-9eE+\-−±×*/().,:%<>=≤≥\s]+", numeric_text)
                if ri > 0 and len(cell_text) > 14 and not numeric:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.1
                paragraph.paragraph_format.space_after = paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.keep_with_next = ri == 0 and bool(rows)
                self.rich(paragraph, cell_text, 9 if count >= 7 else 9.5, bold=ri == 0)
        if block.get("notes"):
            note = self.doc.add_paragraph()
            note.paragraph_format.space_before = Pt(5)
            self.rich(note, block["notes"] if block["notes"].startswith("注") else "注：" + block["notes"], 9.5)
        return caption

    def build(self, data):
        self.doc.core_properties.title = data["title_zh"]
        title = self.doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.rich(title, data["title_zh"], 18, bold=True, heading=True)
        for key, size in (("title_en", 10), ("metadata_line", 9.5)):
            if data.get(key):
                paragraph = self.doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.rich(paragraph, data[key], size)
        for index, block in enumerate(data["blocks"], 1):
            kind = block["kind"]
            if kind == "heading":
                level = block.get("level", 1)
                paragraph = self.doc.add_paragraph(style=f"Heading {level}")
                value = block["zh"]
                if block.get("number") is not None:
                    number = scalar(block["number"], "heading.number")
                    if number and not re.match(rf"^{re.escape(number)}(?:\s|[、.])", value):
                        value = number + "  " + value
                self.rich(paragraph, value, {1: 14, 2: 12, 3: 11}[level], bold=True, heading=True)
            elif kind == "figure":
                paragraph = self.image(block)
                caption = self.doc.add_paragraph(style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.rich(caption, self.numbered("图", block["number"], block["zh"]), 10)
                if block.get("labels_zh"):
                    caption.paragraph_format.keep_with_next = True
                    labels = self.doc.add_paragraph()
                    labels.paragraph_format.space_after = Pt(6)
                    self.rich(labels, "图中术语：" + "；".join(f"{en}：{zh}" for en, zh in block["labels_zh"].items()), 9)
            elif kind == "table":
                paragraph = self.add_table(block)
            elif kind == "equation":
                if block.get("image"):
                    paragraph = self.image(block, equation=True)
                    self.image_equations += 1
                    # A supplied equation crop may already include its original
                    # number. A separate number is added only when requested.
                    if block.get("number") is not None:
                        self.rich(paragraph, "  (" + scalar(block["number"], "equation.number") + ")", 10)
                else:
                    paragraph = self.doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.add_math(paragraph, block["latex"])
                    if block.get("number") is not None:
                        self.rich(paragraph, "  (" + scalar(block["number"], "equation.number") + ")", 10)
            else:
                paragraph = self.doc.add_paragraph()
                if kind == "reference":
                    paragraph.paragraph_format.line_spacing = 1.08
                    paragraph.paragraph_format.space_after = Pt(4)
                    self.rich(paragraph, block["text"], 9)
                elif kind == "list_item":
                    paragraph.paragraph_format.left_indent = Inches(.2)
                    paragraph.paragraph_format.first_line_indent = Inches(-.2)
                    self.rich(paragraph, block["zh"])
                elif kind == "note":
                    self.rich(paragraph, block["zh"], 9.5)
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    self.rich(paragraph, block["zh"])
            if index > 1 and data["blocks"][index - 2]["kind"] == "table":
                paragraph.paragraph_format.space_before = Pt(6)
            self.bookmark(paragraph, block, index)
        return self.doc


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", type=Path, help="Block JSON; image paths are relative to this file")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--mml2omml", type=Path, help="Locally available MML2OMML.XSL (never redistributed)")
    parser.add_argument("--page-size", choices=("letter", "a4"), default="letter")
    parser.add_argument("--font-cjk", default="宋体")
    parser.add_argument("--font-heading-cjk", default="黑体", help="CJK title and heading font (default: 黑体)")
    parser.add_argument("--font-latin", default="Times New Roman")
    parser.add_argument("--report", type=Path, help="Optional structural report; not a visual QA verdict")
    parser.add_argument("--overwrite", action="store_true", help="Explicitly replace existing output/report")
    args = parser.parse_args()
    source, output = args.input.resolve(), args.output.resolve()
    report_path = args.report.resolve() if args.report else None
    if output.suffix.lower() != ".docx":
        parser.error("--output must end in .docx")
    paths = [source, output] + ([report_path] if report_path else [])
    if len(set(paths)) != len(paths):
        parser.error("Input, output and report must be different files")
    for target in [output] + ([report_path] if report_path else []):
        if target.exists() and not args.overwrite:
            parser.error(f"Refusing to overwrite {target}; use a new path or --overwrite")
    try:
        data = json.loads(source.read_text(encoding="utf-8-sig"))
        validate(data)
        builder = Builder(source, args)
        document = builder.build(data)
        if output in builder.asset_paths or report_path in builder.asset_paths:
            raise ValueError("Output/report cannot overwrite an input image")
        output.parent.mkdir(parents=True, exist_ok=True)
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(dir=output.parent, prefix=".docx-build-", suffix=".docx", delete=False) as stream:
                temp_path = Path(stream.name)
            document.save(temp_path)
            if output.exists() and not args.overwrite:
                raise ValueError("Output appeared during build; refusing to overwrite")
            temp_path.replace(output)
        finally:
            if temp_path is not None and temp_path.exists():
                temp_path.unlink()
        report = {"output": str(output), "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
                  "blocks": dict(Counter(block["kind"] for block in data["blocks"])),
                  "native_math_count": builder.math_count, "image_equation_count": builder.image_equations,
                  "front_matter_source_ids": data.get("front_matter_source_ids", {}),
                  "bookmarks": builder.bookmarks, "visual_qa": "NOT_RUN", "semantic_qa": "NOT_RUN"}
        if report_path:
            report_path.parent.mkdir(parents=True, exist_ok=True)
            mode = "w" if args.overwrite else "x"
            with report_path.open(mode, encoding="utf-8") as stream:
                json.dump(report, stream, ensure_ascii=False, indent=2)
                stream.write("\n")
        print(json.dumps(report, ensure_ascii=False))
    except (OSError, ValueError, TypeError, KeyError, etree.Error) as exc:
        parser.exit(1, f"Build failed: {exc}\n")


if __name__ == "__main__":
    main()
